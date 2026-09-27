"""真题文本提取：docx/pypdf 文本层优先，扫描卷回退 OCR 与视觉通道（自 exam_paper_service 拆出）。

提取是"读取"侧的唯一入口：`extract_text` 供扫描探针与导入流水线共用；
PDF 页级提取（`_pdf_pages`/`_pdf_pages_render`）供图示题存原图。
本模块只依赖 scan 的文件识别与 ai_service（视觉通道，函数内局部导入），
不知道导入流水线的存在。
"""

from pathlib import Path
from typing import List

from app.config import settings
from app.services import exam_paper_scan as scan_mod

# AI 拆题前的分块大小：探针估账与流水线拆题共用（原属 exam_paper_service，随使用方迁来）
_CHUNK_CHARS = 9000


def extract_text(path: Path, subject: str = "") -> str:
    """提取 docx / pdf 的全文文本。

    PDF 优先用 pypdf 提取文本层；文本层为空/过少（扫描版图片型 PDF）时，用 pypdfium2
    把页面渲染成图：**数学/408 等公式密集卷优先视觉模型（能输出准确 LaTeX）**，本地
    Windows OCR 作兜底；文科目反之（本地 OCR 快、免费，视觉作兜底），尽量把题面文字捞回。
    """
    suffix = path.suffix.lower()
    if suffix == ".docx":
        import docx

        d = docx.Document(str(path))
        parts = [p.text.strip() for p in d.paragraphs if p.text and p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    if suffix == ".pdf":
        text = _pdf_text_layer(path)
        if _pdf_text_usable(text):
            return text
        # 文本层稀少/为空：扫描版 → 渲染成图，按科目选 视觉/OCR 兜底。
        ocr_text = _pdf_ocr(path, subject)
        if ocr_text.strip():
            return ocr_text
        return text  # 兜底全失败，返回原文本（调用方给出明确报错）
    raise ValueError(f"不支持的文件类型：{suffix}")


def _is_math(subject: str) -> bool:
    """是否为公式密集卷（数学 / 计算机408），以决定优先视觉还是本地 OCR。"""
    s = (subject or "").lower()
    return ("数学" in s) or ("408" in s) or ("计算机" in s)


def _probe_file(path: Path, kind: str) -> dict:
    """导入前"探针"：只花本地 IO 读页数与文本层规模，一次 AI 都不调。

    给 estimate_import 估账单用；依赖缺失（CI 无 pypdf/python-docx）时返回
    ok=False，调用方据此给出"无法预估"的诚实提示而不是编数字。
    """
    out = {"pages": 0, "chars": 0, "usable": False, "ok": False}
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            return out
        try:
            reader = PdfReader(str(path))
            try:
                out["pages"] = len(reader.pages)
                text = "\n".join((p.extract_text() or "") for p in reader.pages)
            finally:
                try:
                    reader.close()
                except Exception:
                    pass
        except Exception:
            return out
        out["chars"] = len(text)
        out["usable"] = _pdf_text_usable(text)
        out["ok"] = True
        return out
    if suffix == ".docx":
        try:
            import docx
        except ImportError:
            return out
        try:
            d = docx.Document(str(path))
            text = "\n".join(p.text for p in d.paragraphs if p.text)
        except Exception:
            return out
        out["chars"] = len(text)
        out["usable"] = bool(text.strip())
        out["ok"] = True
        return out
    return out


def estimate_import(source_rel: str, answer_rel: str = "") -> dict:
    """导入前的"账单"：这份卷要花多少 AI 调用、大约多久，先亮出来用户再拍板。

    全部数字来自本地探针（页数 / 文本层字数），**不调任何 AI**。扫描版、详解册、
    依赖缺失这三类会在 warnings 里明说，宁可说"估不了"也不给假数字。
    """
    root = scan_mod.papers_root()
    source = root / source_rel
    result = {
        "source_path": source_rel,
        "answer_path": answer_rel or "",
        "file_type": source.suffix.lower().lstrip("."),
        "pages": 0,
        "chars": 0,
        "chunks": 0,
        "ai_calls": 0,
        "answer_pages": 0,
        "scanned": False,
        "high_risk": False,
        "estimable": True,
        "warnings": [],
        "minutes": 0,
    }
    if not source.is_file():
        result["estimable"] = False
        result["warnings"].append("源文件不存在，无法预估")
        return result

    role = scan_mod.classify_file(source.name)
    if role == "mixed":
        result["warnings"].append(
            "这是「题+答案合卷/解析册」：页数多、拆题调用成倍多。目录里有纯试题册的话，优先导纯试题册"
        )
    elif role not in ("question",):
        result["warnings"].append(f"文件角色判定为 {role or '未知'}，不是标准试卷")

    probe = _probe_file(source, "source")
    if not probe["ok"]:
        result["estimable"] = False
        result["warnings"].append("服务器缺少 PDF/Word 解析依赖（pypdf / python-docx），无法预估")
        return result

    result["pages"] = probe["pages"]
    result["chars"] = probe["chars"]

    if source.suffix.lower() == ".pdf" and not probe["usable"]:
        result["scanned"] = True
        result["high_risk"] = True
        per_page = min(probe["pages"], int(settings.PDF_OCR_PAGES))
        result["ai_calls"] = per_page
        result["minutes"] = per_page
        result["warnings"].append(
            f"扫描版 PDF（文本层不可用）：要逐页渲染走视觉/OCR，最多 {settings.PDF_OCR_PAGES} 页"
            "——这是最贵最慢的路径，确认前请三思"
        )
    else:
        # 账单口径：按每 _CHUNK_CHARS 字一段估（真实拆题按行切，误差可忽略）
        chunks = max(1, -(-probe["chars"] // _CHUNK_CHARS)) if probe["chars"] else 0
        result["chunks"] = chunks
        result["ai_calls"] = chunks + (1 if answer_rel else 0)
        # 拆题每段按 2~5 分钟估（推理模型 + 12000 token 上限），取中间值
        result["minutes"] = chunks * 3
        if not probe["chars"]:
            result["warnings"].append("未能从文件读到文字，导入会直接失败")
        elif chunks > 6:
            result["warnings"].append(
                f"文本量大（约 {probe['chars'] // 1000} 千字，拆 {chunks} 段），调用次数偏多"
            )

    if answer_rel:
        answer = root / answer_rel
        if not answer.is_file():
            result["warnings"].append("配对的答案文件不存在，导入后答案会全部空缺")
        else:
            ans_probe = _probe_file(answer, "answer")
            result["answer_pages"] = ans_probe["pages"]
            if ans_probe["ok"] and answer.suffix.lower() == ".pdf" and not ans_probe["usable"]:
                # 答案册文本层也不可用时，流水线会强制对它跑整册 OCR（见 _run_import）
                result["high_risk"] = True
                per_page = min(ans_probe["pages"], int(settings.PDF_OCR_PAGES))
                result["ai_calls"] += per_page
                result["minutes"] += per_page
                result["warnings"].append(
                    f"答案册也是扫描版（{ans_probe['pages']} 页）：配答案前要先整册 OCR，再增加约 {per_page} 次调用"
                )
    return result


def _pdf_text_layer(path: Path) -> str:
    """用 pypdf 提取 PDF 文本层（前若干页合并）。"""
    reader = None
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        parts = []
        for page in reader.pages[:80]:
            try:
                parts.append(page.extract_text() or "")
            except Exception:
                continue
        return "\n".join(parts)
    except Exception:
        return ""
    finally:
        # pypdf 6.x 的 PdfReader 持有文件流；不关会让长批次导入累积句柄
        try:
            if reader is not None:
                reader.close()
        except Exception:
            pass


def _pdf_text_usable(text: str) -> bool:
    """文本层是否足够用来拆题。

    扫描版常为空或极稀疏（低于 PDF_TEXT_MIN 即走 OCR）；pypdf 对部分内嵌字体 PDF
    会提取出大量乱码/控制符，故再用「可读字符占比」判定，占比过低同样触发视觉/OCR 兜底。
    """
    s = (text or "").strip()
    if len(s) < settings.PDF_TEXT_MIN:
        return False
    n = len(s)
    cjk = sum(1 for ch in s if "\u4e00" <= ch <= "\u9fff")
    alnum = sum(1 for ch in s if ch.isascii() and ch.isalnum())
    space = sum(1 for ch in s if ch.isspace())
    punct = sum(
        1 for ch in s if ch in "，。、；：？！（）《》【】.,;:?!()[]\"'-+=<>/\\|%$#@&*^~`{}"
    )
    good = cjk + alnum + space + punct
    return good / n >= settings.PDF_TEXT_RATIO


def _pdf_ocr(path: Path, subject: str = "") -> str:
    """扫描版 PDF 兜底：渲染页面为图，逐页提取文字。

    数学/408：视觉模型优先（输出 LaTeX，公式准确），失败退本地 OCR；
    其他科目：本地 Windows OCR 优先（快、免费），失败退视觉模型。
    """
    try:
        import pypdfium2 as pdfium
    except Exception:
        return ""
    from app.services import local_ocr

    try:
        pdf = pdfium.PdfDocument(str(path))
    except Exception:
        return ""
    try:
        return _pdf_ocr_pages(pdf, subject, local_ocr)
    finally:
        # pypdfium2 的文档对象不在 GC 时立刻释放（C 层缓冲），必须显式关，
        # 否则批量导入会一直涨内存 / 占住文件句柄
        try:
            pdf.close()
        except Exception:
            pass


def _pdf_ocr_pages(pdf, subject: str, local_ocr, max_pages: int = 0) -> str:
    import base64
    import io

    pages = len(pdf)
    math = _is_math(subject)
    limit = settings.PDF_OCR_PAGES or pages
    if max_pages:
        limit = min(limit, max_pages)
    out: List[str] = []
    # 正常试卷页数少，全量提取才能覆盖全部题目；仅在累计文本足够拆题(约2段)时提前刹车，
    # 避免 100+ 页的「解析/答案速查」类大文件被整本拖慢。
    needed = _CHUNK_CHARS * 2
    for i in range(min(pages, limit)):
        try:
            page = pdf[i]
            bmp = page.render(scale=2.2)
            pil = bmp.to_pil()
            buf = io.BytesIO()
            pil.save(buf, format="PNG")
            data = buf.getvalue()
        except Exception:
            continue
        b64 = base64.b64encode(data).decode()
        text = ""
        if math:
            # 公式密集：视觉优先（LaTeX），本地 OCR 兜底
            text = _pdf_page_vision(pil, i, subject)
            if not text.strip() and local_ocr.is_available():
                try:
                    text = local_ocr.recognize_base64(b64)
                except Exception:
                    text = ""
        else:
            # 文科目：本地 OCR 优先，视觉兜底
            if local_ocr.is_available():
                try:
                    text = local_ocr.recognize_base64(b64)
                except Exception:
                    text = ""
            if not text.strip():
                text = _pdf_page_vision(pil, i, subject)
        if text.strip():
            # 在每页文本前插入页码标记 [[PAGE:i]]，供 AI 拆题定位该题所在页（图表题用）
            out.append(f"[[PAGE:{i}]]\n{text.strip()}")
        if sum(len(t) for t in out) >= needed:
            break
    return "\n".join(out)


def _pdf_page_vision(pil, index: int, subject: str = "") -> str:
    """单页视觉提取：把渲染图交给 DeepSeek 视觉模型；数学/408 要求输出 LaTeX。"""
    import base64
    import io

    from app.services import ai_service

    buf = io.BytesIO()
    try:
        pil.save(buf, format="PNG")
    except Exception:
        return ""
    b64 = base64.b64encode(buf.getvalue()).decode()
    if _is_math(subject):
        instruction = (
            "这是一页考研数学/计算机408 真题。请完整、准确提取全部文字，"
            "数学公式一律用 LaTeX（如 \\int、\\frac、\\lim、\\sum、\\sqrt、矩阵用 bmatrix/pmatrix）。"
            "只输出内容本身，不要解释。"
        )
    else:
        instruction = (
            "请完整、准确提取这一页试卷的文字，保留段落与换行；如含数学公式请用 LaTeX。"
            "只输出内容本身，不要解释。"
        )
    try:
        return ai_service._vision_extract_text(
            [b64],
            instruction,
            timeout=min(settings.AI_VISION_PRIMARY_TIMEOUT, 90),
        )
    except Exception:
        return ""


def _pdf_pages(path: Path, subject: str) -> list:
    """扫描/公式 PDF 按页提取。返回 [(page_idx, text, pil)]，页码由我直接给定（可靠）。

    数学/408 优先视觉(LaTeX)，本地 OCR 兜底；文科目反之。只保留能提取出文字的页。
    """
    try:
        import pypdfium2 as pdfium
    except Exception:
        return []
    from app.services import local_ocr

    try:
        pdf = pdfium.PdfDocument(str(path))
    except Exception:
        return []
    try:
        return _pdf_pages_render(pdf, subject, local_ocr)
    finally:
        # 同上：必须显式关，否则「渲染 → 视觉调用」的长流程结束前句柄一直占着
        try:
            pdf.close()
        except Exception:
            pass


def _pdf_pages_render(pdf, subject: str, local_ocr) -> list:
    import base64
    import io

    pages = len(pdf)
    math = _is_math(subject)
    limit = settings.PDF_OCR_PAGES or pages
    result: list = []
    for i in range(min(pages, limit)):
        pil = None
        text = ""
        try:
            bmp = pdf[i].render(scale=2.2)
            pil = bmp.to_pil()
            buf = io.BytesIO()
            pil.save(buf, format="PNG")
            b64 = base64.b64encode(buf.getvalue()).decode()
        except Exception:
            continue
        if math:
            text = _pdf_page_vision(pil, i, subject)
            if not text.strip() and local_ocr.is_available():
                try:
                    text = local_ocr.recognize_base64(b64)
                except Exception:
                    text = ""
        else:
            if local_ocr.is_available():
                try:
                    text = local_ocr.recognize_base64(b64)
                except Exception:
                    text = ""
            if not text.strip():
                text = _pdf_page_vision(pil, i, subject)
        if text.strip():
            result.append((i, text.strip(), pil))
    return result
