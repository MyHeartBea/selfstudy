"""真题库：扫描历年真题文件、提取文本、AI 拆题入库（后台流水线）。"""

import json
import queue
import re
import threading
from datetime import datetime, timezone
from pathlib import Path
from typing import List, Optional

from app.config import PROJECT_ROOT, settings
from app.services import ai_service
from app.services.ai_service import AiRequestError

PAPER_STATUSES = ("pending", "extracting", "structuring", "done", "error")

# 科目识别：按 真题库内的路径/文件名 关键词（顺序敏感：408 先于数学）
# 数学要覆盖「数二/数一/数三」这类简写，否则 `2011年数二真题答案速查.pdf` 会识别为空
_SUBJECT_RULES = [
    ("英语二", ("英语二", "英语（二）", "英语")),
    ("计算机408", ("408", "计算机")),
    ("数学二", ("数学二", "数学（二）", "数学", "数二", "数一", "数三", "数学一", "数学三")),
    ("政治", ("政治",)),
]

# 数学一/二/三 不能混：命中「数学（一）/数学一」等就排除在数学二之外
# 注意要覆盖全角括号写法「数学（二）」，否则会一个规则都不命中
_MATH_ONE_KEYS = ("数学（一）", "数学(一)", "数学一", "数一", "数学 一")
_MATH_TWO_KEYS = ("数学（二）", "数学(二)", "数学二", "数二", "数学 二")
_MATH_THREE_KEYS = ("数学（三）", "数学(三)", "数学三", "数三", "数学 三")

# 文件角色判定关键词
# 「答案速查」是明确答案册；「解析/详解」可能是【题+答案合卷】，需要单独判定
_ANSWER_KEYS = ("答案速查", "参考答案", "答案", "解析", "详解")
# 「合卷」特征：文件名同时出现"真题"与"解析/详解"，说明题目和答案在同一份文件里
_MIXED_KEYS = ("解析", "详解")
_PAPER_STEMS = ("真题", "试题", "试卷", "统考")
_CHUNK_CHARS = 9000


def papers_root() -> Path:
    return Path(settings.PAPERS_DIR or str(PROJECT_ROOT / "真题"))


def _classify_math(text: str) -> str:
    """数学卷细分：数学一/二/三 互不混淆（原实现把「2024年数学（一）」当成数学二）。"""
    if any(k in text for k in _MATH_THREE_KEYS):
        return "数学三"
    if any(k in text for k in _MATH_ONE_KEYS):
        return "数学一"
    if any(k in text for k in _MATH_TWO_KEYS):
        return "数学二"
    # 只写了「数学」没写几：按数学二处理（本库主要用数二）
    return "数学二"


def guess_subject(text: str) -> str:
    for subject, keys in _SUBJECT_RULES:
        if any(k in text for k in keys):
            if subject == "数学二":
                return _classify_math(text)
            return subject
    return ""


def guess_year(text: str) -> str:
    """识别年份。优先 4 位年份；否则识别 `26考研`/`25数二` 这类两位年份缩写。"""
    text = str(text or "")
    m = re.search(r"(19|20)\d{2}", text)
    if m:
        return m.group(0)
    # 两位年份：26考研 / 25数二 / 24政治 / 23年
    m = re.search(r"(?<!\d)([0123]\d)\s*(?=考研|数[一二三]|英语|政治|年)", text)
    if m:
        two = int(m.group(1))
        return str(2000 + two) if two <= 40 else str(1900 + two)
    return ""


def _split_compact_year(text: str) -> str:
    """把 `2010-2024` / `2005-2021` 这种区间名里的年份取**开头那个**（用于整体资料文件名）。

    这类文件名通常属于合集，取区间起始年没有意义，故只在没有其它年份时兜底。
    """
    m = re.search(r"(19|20)\d{2}", text)
    return m.group(0) if m else ""


def classify_file(name: str) -> str:
    """判定文件角色：question / answer_key / mixed / other。

    - mixed：**题+答案合卷**（`真题解析`、`真题及参考答案`、`真题+详解`）。以前这类被
      `_ANSWER_KEYS` 一律当成"纯答卷"，于是既当不成试卷也配不到答案，150 份候选里
      52 份因此变成孤儿。
    - answer_key：只是答案/解析册（答案速查、参考答案、选择题解析），不含题面
    - question：纯试卷（真题/试题/试卷/统考，且不含答案解析字样）
    - other：无法判断（答题卡等）

    判定顺序（先具体后笼统，避免"真题答案速查"被当成合卷）：
    1. 含「答案速查」→ 速查答案册；
    2. 含真题/试题/试卷 且 含解析/详解/答案 → 合卷；
    3. 含「参考答案」→ 答案册；
    4. 含解析/详解/答案 → 答案册；
    5. 含真题/试题/试卷 → 纯试卷；
    6. 其它 → other。
    """
    has_stem = any(k in name for k in _PAPER_STEMS)
    has_solution = any(k in name for k in _MIXED_KEYS)
    has_answer = any(k in name for k in _ANSWER_KEYS)

    if "答案速查" in name:
        return "answer_key"
    if has_stem and (has_solution or has_answer):
        return "mixed"
    if "参考答案" in name:
        return "answer_key"
    if has_solution or has_answer:
        return "answer_key"
    if has_stem:
        return "question"
    return "other"


def scan_folder() -> List[dict]:
    """扫描真题根目录，返回**按科目+年份去重后**的候选试卷清单。

    相比旧实现的三处修正：
    1. **角色判定**：区分 纯试卷/答案册/**题+答案合卷**。合卷自己就带答案，
       不再被判成"纯答卷"而变成孤儿（实测 81 份未配对里 52 份是合卷）；
    2. **去重合并**：同一(科目,年份)的「真题」「真题解析」「答案速查」合并成一条候选，
       不再出现英语二某年 3~5 份重复候选、分不清该导哪个；
    3. **年份/科目识别**：补 `26考研→2026` 两位年，数学一/三 不再并入数学二。

    每条候选额外给出 `sources`（该年份涉及的文件与角色）与 `mixed`（是否合卷），
    供前端展示与人工挑选。
    """
    root = papers_root()
    if not root.is_dir():
        return []

    files = []
    for path in sorted(root.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in (".docx", ".pdf"):
            continue
        if path.name.startswith("~$"):
            continue
        rel = str(path.relative_to(root))
        subject = guess_subject(rel)
        year = guess_year(path.name) or guess_year(str(path.parent.name))
        files.append(
            {
                "rel_path": rel,
                "name": path.name,
                "subject": subject,
                "year": year,
                "kind": classify_file(path.name),
                "size_kb": max(1, round(path.stat().st_size / 1024)),
            }
        )

    # 按 (科目, 年份) 归组；年份/科目未识别的单独成组，避免互相污染
    groups: dict = {}
    for f in files:
        if f["kind"] == "other":
            continue
        key = (f["subject"], f["year"])
        groups.setdefault(key, []).append(f)

    results = []
    for (subject, year), items in groups.items():
        questions = [f for f in items if f["kind"] == "question"]
        mixed = [f for f in items if f["kind"] == "mixed"]
        answers = [f for f in items if f["kind"] == "answer_key"]

        # 试卷来源优先级：纯试卷（更像干净题干）> 合卷
        pool = questions or mixed
        if not pool:
            continue
        pool.sort(key=lambda f: (f["kind"] != "question", f["size_kb"]))
        primary = pool[0]

        # 答案来源：明确答案册优先（小文件速查版最干净），其次合卷（合卷里题面会干扰提取）
        answer_keys = answers
        mixed_only = [f for f in mixed if f["rel_path"] != primary["rel_path"]]
        ans_pool = answer_keys or mixed_only
        ans_pool = [f for f in ans_pool if f["rel_path"] != primary["rel_path"]]
        ans_pool.sort(
            key=lambda f: (
                f["kind"] != "answer_key",
                not f["name"].startswith("答案速查"),
                f["size_kb"],
            )
        )
        answer_path = ans_pool[0]["rel_path"] if ans_pool else ""

        results.append(
            {
                "rel_path": primary["rel_path"],
                "name": primary["name"],
                "subject": subject,
                "year": year,
                "kind": "paper",
                "mixed": primary["kind"] == "mixed",
                "size_kb": primary["size_kb"],
                "answer_path": answer_path,
                "answer_kind": ans_pool[0]["kind"] if ans_pool else "",
                "sources": [
                    {"rel_path": f["rel_path"], "kind": f["kind"], "size_kb": f["size_kb"]}
                    for f in sorted(items, key=lambda x: x["rel_path"])
                ],
            }
        )

    results.sort(key=lambda r: (r["subject"], r["year"] == "", r["year"]))
    return results


def extract_text(path: Path, subject: str = "") -> str:
    """提取 docx / pdf 的全文文本。

    PDF 优先用 pypdf 提取文本层；文本层为空/过少（扫描版图片型 PDF）时，用 pypdfium2
    把页面渲染成图：**数学/408 等公式密集卷优先视觉模型（能输出准确 LaTeX）**，本地
    Windows OCR 作兜底；文科目反之（本地 OCR 快、免费，视觉作兜底），尽量把题面文字捞回。
    """
    suffix = path.suffix.lower()
    if suffix == ".docx":
        import docx

        d = docx.Document(str(path))
        parts = [p.text.strip() for p in d.paragraphs if p.text and p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    if suffix == ".pdf":
        text = _pdf_text_layer(path)
        if _pdf_text_usable(text):
            return text
        # 文本层稀少/为空：扫描版 → 渲染成图，按科目选 视觉/OCR 兜底。
        ocr_text = _pdf_ocr(path, subject)
        if ocr_text.strip():
            return ocr_text
        return text  # 兜底全失败，返回原文本（调用方给出明确报错）
    raise ValueError(f"不支持的文件类型：{suffix}")


def _is_math(subject: str) -> bool:
    """是否为公式密集卷（数学 / 计算机408），以决定优先视觉还是本地 OCR。"""
    s = (subject or "").lower()
    return ("数学" in s) or ("408" in s) or ("计算机" in s)


def _pdf_text_layer(path: Path) -> str:
    """用 pypdf 提取 PDF 文本层（前若干页合并）。"""
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
    except Exception:
        return ""
    parts = []
    for page in reader.pages[:80]:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n".join(parts)


def _pdf_text_usable(text: str) -> bool:
    """文本层是否足够用来拆题。

    扫描版常为空或极稀疏（低于 PDF_TEXT_MIN 即走 OCR）；pypdf 对部分内嵌字体 PDF
    会提取出大量乱码/控制符，故再用「可读字符占比」判定，占比过低同样触发视觉/OCR 兜底。
    """
    s = (text or "").strip()
    if len(s) < settings.PDF_TEXT_MIN:
        return False
    n = len(s)
    cjk = sum(1 for ch in s if "\u4e00" <= ch <= "\u9fff")
    alnum = sum(1 for ch in s if ch.isascii() and ch.isalnum())
    space = sum(1 for ch in s if ch.isspace())
    punct = sum(
        1 for ch in s
        if ch in "，。、；：？！（）《》【】.,;:?!()[]\"'-+=<>/\\|%$#@&*^~`{}"
    )
    good = cjk + alnum + space + punct
    return good / n >= settings.PDF_TEXT_RATIO


def _pdf_ocr(path: Path, subject: str = "") -> str:
    """扫描版 PDF 兜底：渲染页面为图，逐页提取文字。

    数学/408：视觉模型优先（输出 LaTeX，公式准确），失败退本地 OCR；
    其他科目：本地 Windows OCR 优先（快、免费），失败退视觉模型。
    """
    import base64
    import io

    try:
        import pypdfium2 as pdfium
    except Exception:
        return ""
    from app.services import local_ocr

    try:
        pdf = pdfium.PdfDocument(str(path))
        pages = len(pdf)
    except Exception:
        return ""

    math = _is_math(subject)
    limit = settings.PDF_OCR_PAGES or pages
    out: List[str] = []
    # 正常试卷页数少，全量提取才能覆盖全部题目；仅在累计文本足够拆题(约2段)时提前刹车，
    # 避免 100+ 页的「解析/答案速查」类大文件被整本拖慢。
    needed = _CHUNK_CHARS * 2
    for i in range(min(pages, limit)):
        try:
            page = pdf[i]
            bmp = page.render(scale=2.2)
            pil = bmp.to_pil()
            buf = io.BytesIO()
            pil.save(buf, format="PNG")
            data = buf.getvalue()
        except Exception:
            continue
        b64 = base64.b64encode(data).decode()
        text = ""
        if math:
            # 公式密集：视觉优先（LaTeX），本地 OCR 兜底
            text = _pdf_page_vision(pil, i, subject)
            if not text.strip() and local_ocr.is_available():
                try:
                    text = local_ocr.recognize_base64(b64)
                except Exception:
                    text = ""
        else:
            # 文科目：本地 OCR 优先，视觉兜底
            if local_ocr.is_available():
                try:
                    text = local_ocr.recognize_base64(b64)
                except Exception:
                    text = ""
            if not text.strip():
                text = _pdf_page_vision(pil, i, subject)
        if text.strip():
            # 在每页文本前插入页码标记 [[PAGE:i]]，供 AI 拆题定位该题所在页（图表题用）
            out.append(f"[[PAGE:{i}]]\n{text.strip()}")
        if sum(len(t) for t in out) >= needed:
            break
    return "\n".join(out)


def _pdf_page_vision(pil, index: int, subject: str = "") -> str:
    """单页视觉提取：把渲染图交给 DeepSeek 视觉模型；数学/408 要求输出 LaTeX。"""
    import base64
    import io

    from app.services import ai_service

    buf = io.BytesIO()
    try:
        pil.save(buf, format="PNG")
    except Exception:
        return ""
    b64 = base64.b64encode(buf.getvalue()).decode()
    if _is_math(subject):
        instruction = (
            "这是一页考研数学/计算机408 真题。请完整、准确提取全部文字，"
            "数学公式一律用 LaTeX（如 \\int、\\frac、\\lim、\\sum、\\sqrt、矩阵用 bmatrix/pmatrix）。"
            "只输出内容本身，不要解释。"
        )
    else:
        instruction = (
            "请完整、准确提取这一页试卷的文字，保留段落与换行；如含数学公式请用 LaTeX。"
            "只输出内容本身，不要解释。"
        )
    try:
        return ai_service._vision_extract_text(
            [b64],
            instruction,
            timeout=min(settings.AI_VISION_PRIMARY_TIMEOUT, 90),
        )
    except Exception:
        return ""


_IMAGE_ROOT = PROJECT_ROOT / "data" / "images" / "exam_papers"
_DIAGRAM_RE = re.compile(r"\[[^\]]*(?:图|示|表|树)[^\]]*\]")
_FIGURE_HINTS = ("页表", "如下表", "如表", "表格如下", "如图所示", "如下图", "示意图", "下图", "如右图")


def _pdf_pages(path: Path, subject: str) -> list:
    """扫描/公式 PDF 按页提取。返回 [(page_idx, text, pil)]，页码由我直接给定（可靠）。

    数学/408 优先视觉(LaTeX)，本地 OCR 兜底；文科目反之。只保留能提取出文字的页。
    """
    import base64
    import io

    try:
        import pypdfium2 as pdfium
    except Exception:
        return []
    from app.services import local_ocr

    try:
        pdf = pdfium.PdfDocument(str(path))
        pages = len(pdf)
    except Exception:
        return []
    math = _is_math(subject)
    limit = settings.PDF_OCR_PAGES or pages
    result: list = []
    for i in range(min(pages, limit)):
        pil = None
        text = ""
        try:
            bmp = pdf[i].render(scale=2.2)
            pil = bmp.to_pil()
            buf = io.BytesIO()
            pil.save(buf, format="PNG")
            b64 = base64.b64encode(buf.getvalue()).decode()
        except Exception:
            continue
        if math:
            text = _pdf_page_vision(pil, i, subject)
            if not text.strip() and local_ocr.is_available():
                try:
                    text = local_ocr.recognize_base64(b64)
                except Exception:
                    text = ""
        else:
            if local_ocr.is_available():
                try:
                    text = local_ocr.recognize_base64(b64)
                except Exception:
                    text = ""
            if not text.strip():
                text = _pdf_page_vision(pil, i, subject)
        if text.strip():
            result.append((i, text.strip(), pil))
    return result


def _paper_img_dir(paper_id: int) -> Path:
    d = _IMAGE_ROOT / str(paper_id)
    d.mkdir(parents=True, exist_ok=True)
    return d


def _has_diagram_option(q: dict) -> bool:
    """题目是否含图示/表格，需要保存该页原图。

    选项为 [图] 类占位，或题干/段落明确引用 页表/表格/如图所示 等，都判为有图。
    """
    opts = " ".join(q.get(k) or "" for k in ("option_a", "option_b", "option_c", "option_d"))
    if _DIAGRAM_RE.search(opts):
        return True
    qtext = (q.get("question") or "") + " " + (q.get("passage") or "")
    return any(h in qtext for h in _FIGURE_HINTS)


def _page_for_no(no: str, page_items: list) -> int:
    """按题号在逐页提取文本里定位该题所在页（可靠，供图示题存图）。

    匹配「行首 题号 + .（、．）」形式，例如 page 文本里 "46.（8 分）..."。
    """
    n = str(no or "").strip()
    if not n.isdigit():
        return 0
    pat = re.compile(r"^[^\S\n]*" + re.escape(n) + r"[.．、]", re.M)
    for i, text, _pil in page_items:
        if pat.search(text or ""):
            return i
    return 0


def _page_diagram_image(paper_id: int, source: Path, page_idx: int, cache: dict, pil=None) -> str:
    """渲染第 page_idx 页为 WebP 图，存 data/images/exam_papers/<pid>/，返回 /images/... URL。

    用于「图示选项」题：把该页截图留存，前端模考/详情可查看原图（同页复用缓存）。
    扫描路径已渲染过该页时，传入 pil 避免重复渲染。
    """
    try:
        key = int(page_idx)
        if key in cache:
            return cache[key]
        import pypdfium2 as pdfium
        from PIL import Image

        if pil is None:
            pdf = pdfium.PdfDocument(str(source))
            if key < 0 or key >= len(pdf):
                return ""
            pil = pdf[key].render(scale=2.4).to_pil()
        pil = pil.convert("RGB")
        if pil.width > 900:
            pil = pil.resize((900, int(pil.height * 900 / pil.width)), Image.LANCZOS)
        out = _paper_img_dir(paper_id) / f"p{key}.webp"
        pil.save(out, format="WEBP", quality=86)
        url = f"/images/exam_papers/{paper_id}/p{key}.webp"
        cache[key] = url
        return url
    except Exception:
        return ""


def _chunk_text(text: str, size: int = _CHUNK_CHARS) -> List[str]:
    chunks = []
    buf = []
    length = 0
    for line in text.splitlines():
        buf.append(line)
        length += len(line) + 1
        if length >= size:
            chunks.append("\n".join(buf))
            buf = []
            length = 0
    if buf:
        chunks.append("\n".join(buf))
    return [c for c in chunks if c.strip()]


def _structure_prompt(subject: str, year: str, chunk: str) -> str:
    return (
        f"你是考研真题整理助手。下面是一份{subject} {year} 年真题试卷的部分文本（可能含考生须知等噪声）。"
        "请把其中的【试题】整理为严格 JSON（不要 Markdown）：\n"
        '{"questions": [{"no": "题号如 1 / 21", "section": "Section I Use of English 等原始节名", '
        '"type": "choice|fill|solution", '
        '"passage": "该题组共用原文（完形填空的文章、阅读理解的全文；只在该题组第一题填写，其他题留空串），无原文则空串", '
        '"question": "题干（选择题为问题句；完形填空为空格所在句；翻译/写作为题目要求全文）", '
        '"option_a": "A 选项", "option_b": "B 选项", "option_c": "C 选项", "option_d": "D 选项", '
        '"analysis": "解析（若文本中带有）", "page": 0, "has_diagram": false}]}\n'
        "规则：\n"
        "1. type 判断：四选项的选 choice；英译汉/翻译与写作选 solution；其余选 fill。\n"
        "2. 只整理试题，跳过考生须知、条形码说明、`[[PAGE:n]]` 页码标记等一切噪声。\n"
        "3. choice 必须带四个选项；选项文本保持原样。**任何数学公式/上下标一律用 `\\(...\\)` 包裹**（如 `\\(2^{8}\\)`、`\\(x^{2}\\)`、`\\(O(n)\\)`），禁止裸写 `^`/`_`。\n"
        "4. correct_answer 一律留空串（答案由系统从答案文件另行匹配）。\n"
        "5. 文本不完整（被截断）时，只整理能完整识别的题目，不要编造。\n"
        "6. 文本中的 `[[PAGE:n]]` 是页码标记：请把该题所在的页码 n 填入 `page` 字段（整数，无标记填 0）。\n"
        "7. 若某选项是图示（树、二叉树、流程图、表格、示意图等，无文字内容），该选项字段填 `[图]`，并置 `has_diagram:true`；"
        "选项有真实文字则如实填写，不要用占位符。\n"
        "8. 请紧凑输出：**每个题目对象单独一行**（不要展开成多行美化），务必是合法 JSON。\n\n"
        f"试卷文本：\n{chunk}"
    )


def _answer_letter(segment: str) -> str:
    """从一段答案文本里取选择题答案字母（A-D），取不到返回空串。"""
    seg = segment[:60]
    for pat in (
        r"^\s*[（(]\s*([A-Da-d])\s*[)）]",   # (A) / （A）
        r"^\s*([A-Da-d])\s*[.、．)）]",       # A. / A、
        r"^\s*([A-Da-d])(?![A-Za-z])",        # 裸 A
    ):
        m = re.match(pat, seg)
        if m:
            return m.group(1).upper()
    return ""


# 一、选择题:1～10 小题 …（用题干里的题号区间关联答案，而不是答案文件自己的序号）
# 注意分隔符里 `-` 必须放最后或转义，否则会被当成范围（`[～~\-—至]` 里 `~-—` 是范围
# 且不含 `-`，曾导致 "17-22题" 匹配失败）
_QUESTION_RANGE_RE = re.compile(
    r"([一二三四五六七八九十]+)\s*[、.．]?\s*(选择题|填空题|解答题|单项选择|多项选择)"
    r"[^\n\d]{0,30}?(\d{1,2})\s*[～~—至\-–]\s*(\d{1,2})"
)
# 题型关键字 → 统一题型
_SECTION_TYPE = {
    "选择题": "choice",
    "单项选择": "choice",
    "填空题": "fill",
    "多项选择": "choice",
    "解答题": "solution",
}


def _answer_section_map(exam_text: str) -> dict:
    """从**试卷**文本里解析「题号 → 题型」映射。

    答案文件里的 `1～10` 是答案册自己的序号，而卷面题号可能是连续的
    （如选择题 1-10，解答题 17-22），不能按序号硬套；用它来校正题型。
    """
    mapping: dict = {}
    for m in _QUESTION_RANGE_RE.finditer(exam_text or ""):
        qtype = _SECTION_TYPE.get(m.group(2))
        if not qtype:
            continue
        try:
            start, end = int(m.group(3)), int(m.group(4))
        except ValueError:
            continue
        if end < start or end - start > 60:
            continue
        for n in range(start, end + 1):
            mapping.setdefault(str(n), qtype)
    return mapping


def _normalize_answer_text(text: str) -> str:
    """把各种"答案册"写法归一成 `题号:答案` 行，大幅提高答案匹配成功率。

    实测过会失败的两种写法（此前 10 道选择题只配到 1 道，而且配错）：
      1) 连排速查式：`(1)C. (2)B. (3)C. …`        → 1:C 2:B 3:C …
      2) 逐题式：    `1【答案】（A）$a=\\frac{1}{3}$ 考点：泰勒公式`
                    → 1:A（选择题只取字母，后面的解析不参与匹配）
    已经是 `1.A` / `1:A` 这类写法的则原样保留。
    """
    if not text or not text.strip():
        return text

    out_lines = []
    for line in text.split("\n"):
        stripped = line.strip()
        if not stripped:
            out_lines.append(line)
            continue

        # 1) 连排速查：(1)C. (2)B. → 拆成多行
        pairs = re.findall(
            r"[（(]?\s*(\d{1,2})\s*[)）.、．:：]\s*[（(]?\s*([A-Da-d])\s*[)）.、．]?",
            stripped,
        )
        if len(pairs) >= 3:
            for no, letter in pairs:
                out_lines.append(f"{int(no)}:{letter.upper()}")
            continue

        # 2) 逐题式：1【答案】（A）/ 1.【答案】A / 11【答案】$0<p<2$
        m = re.match(r"^\s*(\d{1,2})\s*[.、．]?\s*[【\[]\s*答案\s*[】\]]\s*(.+)$", stripped)
        if m:
            no, rest = m.group(1), m.group(2).strip()
            letter = _answer_letter(rest)
            # 判断"这是选择题字母答案还是数学表达式"：
            # 只看开头 —— `（A）$a=\frac{1}{3}$` 里的 `$` 出现在字母之后，
            # 不能因此否认它是选择题答案（这是之前 1:A 提取不到的根因）
            is_expression = re.match(r"^\s*(\\|\$\$|[a-zA-Z]\s*[=<>])", rest)
            if letter and not is_expression:
                out_lines.append(f"{int(no)}:{letter}")
            else:
                out_lines.append(f"{int(no)}:{rest}")
            continue

        out_lines.append(line)

    return "\n".join(out_lines)


def _answers_prompt(subject: str, year: str, nos: List[str], answer_text: str) -> str:
    return (
        f"以下是{subject} {year} 年真题的答案材料。\n"
        "材料可能已经过预处理，形如 `题号:答案`（如 `3:C`）；也可能格式凌乱。\n\n"
        "任务：为下面列出的每个题号给出正确答案，输出严格 JSON："
        '{"answers": {"题号": "A/B/C/D"}}\n\n'
        "**准确优先，宁缺勿错**（重要）：\n"
        "1. 逐条核对材料里的题号，只填你能在材料中找到明确依据的题；\n"
        "2. 材料里找不到的题号**必须省略**，绝对不要靠推理、经验或“看起来像”来猜；\n"
        "3. 材料里若出现题号重复或互相矛盾，以第一次出现为准，仍不确定就省略；\n"
        "4. value 只填单个大写字母 A/B/C/D，不要带括号、句点或中文解释。\n\n"
        f"需要匹配的题号：{json.dumps(nos, ensure_ascii=False)}\n\n"
        f"答案材料：\n{answer_text[:12000]}"
    )


# —— 后台导入流水线：单工作线程串行消费，避免并发 AI 互相踩 ——
_import_queue: "queue.Queue[int]" = queue.Queue()
_worker_lock = threading.Lock()
_worker_started = False


def enqueue_import(paper_id: int) -> None:
    global _worker_started
    _import_queue.put(paper_id)
    with _worker_lock:
        if not _worker_started:
            _worker_started = True
            t = threading.Thread(target=_worker_loop, daemon=True, name="km-paper-import")
            t.start()


def _set_status(conn, paper_id: int, status: str, note: str = "") -> None:
    conn.execute(
        "UPDATE exam_papers SET status = ?, status_note = ? WHERE id = ?",
        (status, note[:300], paper_id),
    )
    conn.commit()


def _worker_loop() -> None:
    from app.database import get_connection

    while True:
        paper_id = _import_queue.get()
        try:
            _run_import(paper_id)
        except Exception as exc:  # 兜底：任何异常都落为 error 状态
            conn = get_connection()
            try:
                _set_status(conn, paper_id, "error", str(exc))
            finally:
                conn.close()


def _run_import(paper_id: int) -> None:
    from app.database import get_connection

    conn = get_connection()
    try:
        row = conn.execute("SELECT * FROM exam_papers WHERE id = ?", (paper_id,)).fetchone()
        if row is None:
            return
        paper = dict(row)
        root = papers_root()
        source = root / paper["source_path"]
        if not source.is_file():
            _set_status(conn, paper_id, "error", f"源文件不存在：{paper['source_path']}")
            return

        _set_status(conn, paper_id, "extracting", "正在提取试卷文本")
        page_pils: dict = {}
        if source.suffix.lower() == ".docx":
            # docx 没有"文本层/扫描版"之分：直接抽全文。
            # 原实现无条件调用 _pdf_text_layer（只认 PDF），docx 会得到空串 →
            # 被误判成"扫描版" → 白跑一轮视觉/OCR → 失败时还报出错误的
            # "扫描版且 OCR/视觉均失败"（实测 25考研政治真题.docx、2024考研英语二真题.docx 都栽在这）。
            pages = []
            scanned = False
            try:
                text_layer = extract_text(source, paper["subject"])
            except Exception as exc:
                _set_status(conn, paper_id, "error", f"docx 读取失败：{exc}")
                return
        else:
            text_layer = _pdf_text_layer(source)
            scanned = not _pdf_text_usable(text_layer)
            pages = _pdf_pages(source, paper["subject"]) if scanned else []
        if scanned and not pages:
            _set_status(conn, paper_id, "error", "未能从文件提取到文本（扫描版且 OCR/视觉均失败）")
            return
        exam_text = text_layer if not scanned else ("\n".join(t for _, t, _ in pages) or text_layer)
        if not exam_text.strip():
            _set_status(conn, paper_id, "error", "未能从文件提取到文本（可能是扫描版 PDF，请换 Word/文本版）")
            return

        answer_text = ""
        if paper["answer_path"]:
            answer_path = root / paper["answer_path"]
            if answer_path.is_file():
                try:
                    answer_text = extract_text(answer_path)
                except Exception:
                    answer_text = ""
        # 答案册写法五花八门（连排速查 `(1)C. (2)B.`、逐题 `1【答案】（A）…考点：…`），
        # 先归一成 `题号:答案` 行再交给 AI，否则匹配率极低（实测 10 题只配到 1 题且配错）
        if answer_text:
            answer_text = _normalize_answer_text(answer_text)
        # 从卷面解析题号→题型，用于校正答案册序号与卷面题号不一致的情况
        answer_section_map = _answer_section_map(exam_text)

        _set_status(conn, paper_id, "structuring", "AI 正在拆题")
        questions: List[dict] = []
        seen_nos = set()

        def _collect(raw_questions, page_idx: int | None = None) -> None:
            for q in raw_questions:
                no = str(q.get("no") or "").strip()
                qtype = str(q.get("type") or "choice")
                if qtype not in ("choice", "fill", "solution"):
                    qtype = "choice"
                if not no or not str(q.get("question") or "").strip():
                    continue
                key = f"{no}|{qtype}"
                if key in seen_nos:
                    continue
                seen_nos.add(key)
                q_page = page_idx if page_idx is not None else int(q.get("page") or 0)
                questions.append(
                    {
                        "no": no,
                        "section": str(q.get("section") or "")[:80],
                        "type": qtype,
                        "passage": str(q.get("passage") or ""),
                        "question": str(q.get("question") or ""),
                        "option_a": str(q.get("option_a") or ""),
                        "option_b": str(q.get("option_b") or ""),
                        "option_c": str(q.get("option_c") or ""),
                        "option_d": str(q.get("option_d") or ""),
                        "correct_answer": "",
                        "analysis": str(q.get("analysis") or ""),
                        "page_idx": q_page,
                        "diagram_image": "",
                    }
                )

        if scanned:
            # 统一扁平分块拆题（对扫描/公式卷更完整，避免逐页漏掉同页多个综合题）；
            # 拆完再按题号在逐页文本里定位页码（可靠，供图示题存该页原图）。
            for i, page_text, pil in pages:
                page_pils[i] = pil
            chunks = _chunk_text(exam_text)
            for idx, chunk in enumerate(chunks):
                parsed = ai_service._chat_json(
                    [{"role": "user", "content": _structure_prompt(paper["subject"], paper["year"], chunk)}],
                    max_tokens=12000,
                )
                _collect(parsed.get("questions", []) or [])
                _set_status(conn, paper_id, "structuring", f"AI 拆题中 {idx + 1}/{len(chunks)} 段")
            for q in questions:
                q["page_idx"] = _page_for_no(q["no"], pages)
        else:
            chunks = _chunk_text(exam_text)
            for idx, chunk in enumerate(chunks):
                parsed = ai_service._chat_json(
                    [{"role": "user", "content": _structure_prompt(paper["subject"], paper["year"], chunk)}],
                    max_tokens=12000,
                )
                _collect(parsed.get("questions", []) or [])
                _set_status(conn, paper_id, "structuring", f"AI 拆题中 {idx + 1}/{len(chunks)} 段")

        if not questions:
            _set_status(conn, paper_id, "error", "AI 未能从文本中整理出试题")
            return

        # 用卷面的「题型-题号区间」校正题型：AI 有时把解答题也标成 choice，
        # 那会让答案匹配去问根本不存在的选项答案（也影响模考只取客观题）。
        if answer_section_map:
            fixed = 0
            for q in questions:
                expect = answer_section_map.get(str(q["no"]))
                if expect and q["type"] != expect:
                    q["type"] = expect
                    fixed += 1
                    if expect != "choice":
                        q["correct_answer"] = ""  # 主观题不留选择题式答案
            if fixed:
                _set_status(conn, paper_id, "structuring", f"已按卷面校正 {fixed} 道题的题型")

        # 答案匹配：仅客观题，从配对答案文件文本推断
        if answer_text:
            need = [q["no"] for q in questions if q["type"] == "choice" and not q["correct_answer"]]
            if need:
                _set_status(conn, paper_id, "structuring", "正在匹配参考答案")
                try:
                    parsed = ai_service._chat_json(
                        [
                            {
                                "role": "user",
                                "content": _answers_prompt(
                                    paper["subject"], paper["year"], need, answer_text
                                ),
                            }
                        ],
                        max_tokens=4000,
                    )
                    answers = parsed.get("answers") or {}
                    for q in questions:
                        if q["type"] == "choice" and q["no"] in answers:
                            ans = str(answers[q["no"]] or "").strip().upper()
                            if re.fullmatch(r"[A-D]", ans):
                                q["correct_answer"] = ans
                except (AiRequestError, Exception):
                    pass  # 答案匹配失败不阻塞入库，答案可后续补

        # 图示选项题：保存该页题图（树/图/流程等原图），供模考/详情查看
        if any(_has_diagram_option(q) for q in questions):
            _img_cache: dict = {}
            for q in questions:
                if scanned and _has_diagram_option(q):
                    page = q.get("page_idx", 0)
                    q["diagram_image"] = _page_diagram_image(
                        paper_id, source, page, _img_cache, page_pils.get(page)
                    )

        conn.execute("DELETE FROM exam_questions WHERE paper_id = ?", (paper_id,))
        now_text = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
        for q in questions:
            conn.execute(
                "INSERT INTO exam_questions (paper_id, no, section, question_type, passage, "
                "question, option_a, option_b, option_c, option_d, correct_answer, analysis, "
                "page_idx, diagram_image) "
                "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (
                    paper_id,
                    q["no"],
                    q["section"],
                    q["type"],
                    q["passage"],
                    q["question"],
                    q["option_a"],
                    q["option_b"],
                    q["option_c"],
                    q["option_d"],
                    q["correct_answer"],
                    q["analysis"],
                    q.get("page_idx", 0),
                    q.get("diagram_image", ""),
                ),
            )
        answered = sum(1 for q in questions if q["type"] == "choice" and q["correct_answer"])
        conn.execute(
            "UPDATE exam_papers SET status = 'done', status_note = ?, question_count = ? WHERE id = ?",
            (f"客观题已配答案 {answered}/{sum(1 for q in questions if q['type'] == 'choice')}", len(questions), paper_id),
        )
        conn.commit()
    finally:
        conn.close()
